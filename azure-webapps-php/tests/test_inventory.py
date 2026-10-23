from io import StringIO
import json
from pathlib import Path
import sys
import tempfile
import unittest
from contextlib import redirect_stderr

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import inventory as inv

SID = "11111111-1111-1111-1111-111111111111"
BASE = f"/subscriptions/{SID}/resourceGroups/test/providers/Microsoft.Web/sites/"
SUB = {"id": SID, "name": "Synthetic subscription", "state": "Enabled"}


def site(name="php-app", pna="Enabled", kind="app,linux"):
    return {"id": BASE + name, "name": name, "kind": kind, "location": "westeurope",
            "properties": {"state": "Running", "publicNetworkAccess": pna,
                           "defaultHostName": name + ".azurewebsites.net", "httpsOnly": True}}


def rule(ip="Any", action="Allow", priority=100, **extra):
    return {"ipAddress": ip, "action": action, "priority": priority, **extra}


class RuntimeTests(unittest.TestCase):
    def test_linux_php_branch_does_not_claim_patch_level(self):
        result = inv.php_runtime({"linuxFxVersion": "PHP|8.2"}, "app,linux")
        self.assertEqual(result["version"], "8.2")
        self.assertEqual(result["evidence"], "PHP configurado")

    def test_linux_node_ignores_stale_php_version(self):
        self.assertIsNone(inv.php_runtime({"linuxFxVersion": "NODE|22-lts", "phpVersion": "8.2"}, "app,linux"))

    def test_windows_php_and_off(self):
        self.assertEqual(inv.php_runtime({"phpVersion": "8.2"}, "app")["source"], "phpVersion")
        for value in (None, "", "Off", "0"):
            self.assertIsNone(inv.php_runtime({"phpVersion": value}, "app"))

    def test_container_is_not_installed_evidence(self):
        config = {"linuxFxVersion": "DOCKER|php:8.2-fpm", "phpVersion": "8.2"}
        self.assertIsNone(inv.php_runtime(config, "app,linux,container"))
        hint = inv.php_runtime(config, "app,linux,container", True)
        self.assertIn("NO confirmado", hint["evidence"])
        self.assertNotIn("php:8.2-fpm", hint["runtime"])
        self.assertIsNone(inv.php_runtime({"linuxFxVersion": "DOCKER|example.invalid/custom:latest"}, "app,linux", True))

    def test_functions_are_excluded(self):
        self.assertFalse(inv.web_app({"kind": "functionapp,linux"}))
        self.assertTrue(inv.web_app({"kind": "app,linux"}))


class RestrictionTests(unittest.TestCase):
    def assess(self, rules=None, default=None):
        config = {"ipSecurityRestrictions": rules or []}
        if default:
            config["ipSecurityRestrictionsDefaultAction"] = default
        return inv.evaluate_rules(config)

    def test_empty_rules_are_public_by_default(self):
        self.assertEqual(self.assess()["exposure"], "Sí")

    def test_empty_rules_explicit_default_deny(self):
        self.assertEqual(self.assess(default="Deny")["exposure"], "No")

    def test_allow_public_cidr_implies_default_deny_and_restricted(self):
        result = self.assess([rule("8.8.8.0/24")])
        self.assertEqual(result["exposure"], "Restringida")
        self.assertEqual(result["default"], "Deny")

    def test_allow_all_ipv4_is_open_even_with_ipv6_denied(self):
        result = self.assess([rule("0.0.0.0/0")])
        self.assertEqual(result["exposure"], "Sí")
        self.assertEqual(result["ipv4_open"], "Sí")

    def test_deny_all_before_allow_all(self):
        self.assertEqual(self.assess([rule("Any", "Deny", 10), rule("Any", "Allow", 20)])["exposure"], "No")

    def test_allow_range_before_deny_all(self):
        self.assertEqual(self.assess([rule("8.8.8.0/24", "Allow", 10), rule("Any", "Deny", 20)])["exposure"], "Restringida")

    def test_earlier_deny_covers_later_allow(self):
        self.assertEqual(self.assess([rule("8.8.0.0/16", "Deny", 10), rule("8.8.8.0/24", "Allow", 20)])["exposure"], "No")

    def test_deny_half_then_allow_other_half_not_unrestricted(self):
        self.assertEqual(self.assess([rule("8.8.8.0/25", "Deny", 10), rule("8.8.8.0/24", "Allow", 20)])["exposure"], "Restringida")

    def test_private_ranges_or_vnet_only_are_not_public(self):
        self.assertEqual(self.assess([rule("10.0.0.0/8")])["exposure"], "No")
        self.assertEqual(self.assess([rule("AnyVnets")])["exposure"], "No")
        self.assertEqual(self.assess([rule("", vnetSubnetResourceId="/vnet/subnet")])["exposure"], "No")

    def test_frontdoor_tag_and_header_are_restricted_not_private(self):
        config = [rule("AzureFrontDoor.Backend", tag="ServiceTag", headers={"x-azure-fdid": ["synthetic"]})]
        self.assertEqual(self.assess(config)["exposure"], "Restringida")

    def test_header_only_allow_does_not_mean_open(self):
        self.assertEqual(self.assess([rule(headers={"x-forwarded-host": ["example.test"]})])["exposure"], "Restringida")

    def test_malformed_rule_and_priority_ties_are_unknown(self):
        self.assertEqual(self.assess([rule("not-an-address")])["exposure"], "Indeterminada")
        self.assertEqual(self.assess([rule("Any", "Allow", 10), rule("Any", "Deny", 10)])["exposure"], "Indeterminada")

    def test_ipv6_public_rule(self):
        self.assertEqual(self.assess([rule("2606:4700::/32")])["exposure"], "Restringida")


class NetworkTests(unittest.TestCase):
    def test_pna_disabled_is_private_even_without_config(self):
        self.assertEqual(inv.network_assessment(site(pna="Disabled"), {}, False, None, None)["exposure"], "No")

    def test_private_endpoint_does_not_override_explicit_public_enable(self):
        result = inv.network_assessment(site(), {}, True, [{"properties": {}}], None)
        self.assertEqual(result["exposure"], "Sí")

    def test_unset_pna_defaults_depend_on_private_endpoints(self):
        app = site(pna=None)
        self.assertEqual(inv.network_assessment(app, {}, True, [], None)["exposure"], "Sí")
        self.assertEqual(inv.network_assessment(app, {}, True, [{}], None)["exposure"], "No")
        self.assertEqual(inv.network_assessment(app, {}, True, None, None)["exposure"], "Indeterminada")

    def test_missing_config_is_not_treated_as_empty_rules(self):
        self.assertEqual(inv.network_assessment(site(), {}, False, [], None)["exposure"], "Indeterminada")

    def test_vnet_integration_is_not_inbound_privacy(self):
        app = site()
        app["properties"]["virtualNetworkSubnetId"] = "/outbound-integration"
        self.assertEqual(inv.network_assessment(app, {}, True, [], None)["exposure"], "Sí")

    def test_internal_ase_unknown_ase_and_publishing_only(self):
        app = site()
        app["properties"]["hostingEnvironmentProfile"] = {"id": "/ase"}
        self.assertEqual(inv.network_assessment(app, {}, True, [], {"properties": {"internalLoadBalancingMode": "Web, Publishing"}})["exposure"], "No")
        self.assertEqual(inv.network_assessment(app, {}, True, [], None)["exposure"], "Indeterminada")
        self.assertEqual(inv.network_assessment(app, {}, True, [], {"properties": {"internalLoadBalancingMode": "Publishing"}})["exposure"], "Sí")


class FakeClient:
    def __init__(self):
        self.php = site()
        self.node = site("node-app")
        self.slot = site("node-app/staging")
        self.slot["id"] = self.node["id"] + "/slots/staging"
        self.node["kind"] = "app,linux"
        self.failure = None
        self.reads = []

    def pages(self, path):
        self.reads.append(path)
        if path.startswith(f"/subscriptions/{SID}/providers/"):
            yield self.php
            yield self.node
        elif "/slots?" in path:
            if path.startswith(self.node["id"] + "/slots?"):
                yield self.slot
        elif "/privateEndpointConnections?" in path:
            if self.failure == "private":
                raise inv.ReadError("AuthorizationFailed", "Synthetic read denial")
        else:
            raise AssertionError(path)

    def get(self, path):
        self.reads.append(path)
        if self.failure == "config" and path.startswith(self.php["id"] + "/config/"):
            raise inv.ReadError("AuthorizationFailed", "Synthetic read denial")
        version = "NODE|22-lts" if path.startswith(self.node["id"] + "/config/") else "PHP|8.2"
        return {"properties": {"linuxFxVersion": version,
                               "appSettings": [{"name": "PASSWORD", "value": "DO_NOT_EXPORT_SYNTHETIC_SECRET"}],
                               "ipSecurityRestrictions": []}}


class InventoryTests(unittest.TestCase):
    def collect(self, client, **kwargs):
        with redirect_stderr(StringIO()):
            return inv.inventory(client, [SUB], workers=1, **kwargs)

    def test_php_slot_of_non_php_parent_is_included(self):
        client = FakeClient()
        data = self.collect(client)
        self.assertEqual(len(data["php"]), 2)
        self.assertEqual(data["scanned"], 3)
        self.assertEqual({r["Tipo"] for r in data["php"]}, {"Web App", "Slot"})
        self.assertNotIn("DO_NOT_EXPORT", json.dumps(data))
        self.assertNotIn("appSettings", json.dumps(data))
        self.assertTrue(all("listPublishing" not in p and "appsettings" not in p for p in client.reads))

    def test_partial_config_denial_keeps_fallback_as_unknown(self):
        client = FakeClient()
        client.php["properties"]["siteConfig"] = {"linuxFxVersion": "PHP|8.2"}
        client.failure = "config"
        data = self.collect(client)
        row = next(r for r in data["php"] if r["Tipo"] == "Web App")
        self.assertEqual(row["Exposición pública configurada"], "Indeterminada")
        self.assertTrue(data["errors"])

    def test_private_endpoint_denial_keeps_unknown_default(self):
        client = FakeClient()
        client.php["properties"]["publicNetworkAccess"] = None
        client.failure = "private"
        data = self.collect(client)
        row = next(r for r in data["php"] if r["Tipo"] == "Web App")
        self.assertEqual(row["Exposición pública configurada"], "Indeterminada")

    def test_later_page_failure_preserves_earlier_resources(self):
        client = FakeClient()
        original = client.pages
        def pages(path):
            if path.startswith(f"/subscriptions/{SID}/providers/"):
                yield client.php
                raise inv.ReadError("AuthorizationFailed", "Synthetic later-page denial")
            yield from original(path)
        client.pages = pages
        data = self.collect(client)
        self.assertEqual(len(data["php"]), 1)
        self.assertEqual(data["coverage"][0]["Enumeración completa"], "No")

    def test_excel_is_filtered_and_does_not_contain_formulas_or_secrets(self):
        from openpyxl import load_workbook
        data = self.collect(FakeClient())
        data["php"][0]["Web App / slot"] = "=HYPERLINK(\"https://example.test\")"
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp) / "inventory.xlsx"
            inv.write_excel(data, out)
            wb = load_workbook(out)
            self.assertEqual(wb.sheetnames, ["Resumen", "WebApps PHP", "Reglas acceso", "Incidencias", "Cobertura"])
            self.assertEqual(wb["WebApps PHP"].max_row, 3)
            self.assertEqual(wb["WebApps PHP"]["D2"].data_type, "s")
            self.assertFalse(any(c.data_type == "f" for ws in wb for row in ws for c in row))
            self.assertFalse(any("DO_NOT_EXPORT" in str(c.value) for ws in wb for row in ws for c in row))
            self.assertEqual(wb["WebApps PHP"].freeze_panes, "E2")
            with self.assertRaises(inv.ReadError):
                inv.write_excel(data, out)

    def test_azure_pagination_follows_trusted_arm_only(self):
        client = inv.AzureCLI.__new__(inv.AzureCLI)
        client.endpoint = "https://management.azure.com"
        calls = []
        responses = iter([{"value": [{"id": "first"}], "nextLink": "https://management.azure.com/next"}, {"value": [{"id": "second"}]}])
        def command(args):
            calls.append(args)
            return next(responses)
        client.command = command
        self.assertEqual([r["id"] for r in client.pages("/first")], ["first", "second"])
        self.assertTrue(all(a[:3] == ["rest", "--method", "get"] for a in calls))
        with self.assertRaises(inv.ReadError):
            client.get("https://other.test/steal")

    def test_rest_selects_subscription_without_account_set(self):
        client = inv.AzureCLI.__new__(inv.AzureCLI)
        client.endpoint = "https://management.azure.com"
        calls = []
        client.command = lambda args: calls.append(args) or {"properties": {}}
        client.get(BASE + "php-app/config/web?api-version=" + inv.API)
        self.assertEqual(calls[0][-2:], ["--subscription", SID])
        self.assertNotIn("set", calls[0])


if __name__ == "__main__":
    unittest.main()
