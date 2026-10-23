#!/usr/bin/env python3
"""Read-only Azure Web Apps PHP/network inventory. Requires an existing az login."""
from __future__ import annotations

import argparse
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
import ipaddress
import json
from pathlib import Path
import re
import shutil
import subprocess
import sys
import time
from urllib.parse import urlsplit
from uuid import UUID

API = "2024-11-01"
NETWORK_DOC = "https://learn.microsoft.com/azure/app-service/overview-access-restrictions"


class ReadError(RuntimeError):
    def __init__(self, code, message):
        super().__init__(message)
        self.code = code


class AzureCLI:
    """Never requests tokens, app settings, publishing credentials or writes to Azure."""
    def __init__(self, timeout=60):
        self.timeout = timeout
        self.executable = shutil.which("az") or shutil.which("az.cmd")
        if not self.executable:
            raise ReadError("AzureCLINotFound", "Instala Azure CLI o usa Azure Cloud Shell (Bash).")
        cloud = self.command(["cloud", "show", "--output", "json"])
        self.endpoint = cloud["endpoints"]["resourceManager"].rstrip("/")
        if urlsplit(self.endpoint).scheme != "https":
            raise ReadError("InvalidCloud", "El endpoint ARM debe utilizar HTTPS.")

    def command(self, args):
        try:
            proc = subprocess.run([self.executable, *args, "--only-show-errors"],
                                  capture_output=True, text=True, timeout=self.timeout,
                                  encoding="utf-8", errors="replace", shell=False)
        except subprocess.TimeoutExpired as exc:
            raise ReadError("Timeout", "Azure CLI superó el tiempo máximo de lectura.") from exc
        if proc.returncode:
            # stderr can contain URLs or sensitive configuration; never persist it verbatim.
            text = proc.stderr
            match = re.search(r"\b(AuthorizationFailed|Forbidden|TooManyRequests|NotFound|ResourceNotFound|InvalidAuthenticationToken|SubscriptionNotFound|ExpiredAuthenticationToken|ServiceUnavailable|InternalServerError)\b", text)
            code = match.group(1) if match else "AzureCLIError"
            if "az login" in text.lower() or "please run" in text.lower():
                code = "LoginRequired"
            raise ReadError(code, "La lectura de Azure falló; revisa sesión, alcance RBAC y disponibilidad.")
        try:
            return json.loads(proc.stdout)
        except json.JSONDecodeError as exc:
            raise ReadError("InvalidJSON", "Azure CLI no devolvió JSON válido.") from exc

    def get(self, path):
        url = path if path.startswith("https://") else self.endpoint + path
        parsed, root = urlsplit(url), urlsplit(self.endpoint)
        if parsed.scheme != "https" or parsed.netloc != root.netloc or parsed.username:
            raise ReadError("InvalidNextLink", "Se rechazó un endpoint ajeno al ARM de la nube seleccionada.")
        for attempt in range(3):
            try:
                args = ["rest", "--method", "get", "--url", url, "--output", "json"]
                subscription = re.search(r"/subscriptions/([0-9a-fA-F-]{36})(?:/|\?|$)", parsed.path + "?", re.I)
                if subscription:
                    # Select credentials per request; never mutate the user's default subscription.
                    args.extend(["--subscription", subscription.group(1)])
                return self.command(args)
            except ReadError as exc:
                if exc.code not in {"TooManyRequests", "ServiceUnavailable", "InternalServerError"} or attempt == 2:
                    raise
                time.sleep(2 ** attempt)

    def pages(self, path):
        # Preserve each successful page if a subsequent page fails.
        seen = set()
        while path:
            if path in seen:
                raise ReadError("PaginationLoop", "Azure devolvió un nextLink repetido.")
            seen.add(path)
            page = self.get(path)
            if not isinstance(page.get("value"), list):
                raise ReadError("InvalidPage", "La lista ARM no contiene value como colección.")
            yield from page["value"]
            path = page.get("nextLink")

    def subscriptions(self):
        return self.command(["account", "list", "--all", "--output", "json"])


def properties(resource):
    value = resource.get("properties")
    return value if isinstance(value, dict) else {}


def web_app(resource):
    kinds = set(str(resource.get("kind", "")).lower().split(","))
    return not kinds.intersection({"functionapp", "workflowapp", "logicapp"})


def php_runtime(config, kind="", include_image_hints=False):
    """Returns configured-runtime evidence, not an assertion that the process is running."""
    linux = str(config.get("linuxFxVersion") or "").strip()
    windows = str(config.get("windowsFxVersion") or "").strip()
    fx = linux or windows
    if re.match(r"^PHP\|", fx, re.I):
        return {"version": fx.split("|", 1)[1], "evidence": "PHP configurado", "source": "linuxFxVersion" if linux else "windowsFxVersion", "runtime": fx}
    if fx and not re.match(r"^(DOCKER|COMPOSE|KUBE)\||^sitecontainers$", fx, re.I):
        return None  # e.g. Linux NODE|22 must not be labelled PHP by a stale phpVersion.
    container = bool(re.match(r"^(DOCKER|COMPOSE|KUBE)\||^sitecontainers$", fx, re.I)) or "container" in kind.lower()
    if container:
        if include_image_hints and re.match(r"^DOCKER\|", fx, re.I):
            image = fx.split("|", 1)[1]
            # A name is a hint only; never infer the contents from a generic custom image.
            if re.search(r"(?:^|[/_.-])php(?:[:/_.-]|$)", image, re.I):
                return {"version": "No verificable por el nombre de imagen", "evidence": "Indicio de imagen; NO confirmado", "source": "Nombre de imagen Docker", "runtime": "DOCKER (nombre no exportado)"}
        return None
    version = str(config.get("phpVersion") or "").strip()
    if "linux" not in kind.lower() and version.lower() not in {"", "off", "disabled", "none", "0", "false"}:
        return {"version": version, "evidence": "PHP configurado", "source": "phpVersion", "runtime": "Windows PHP " + version}
    return None


def _subtract(ranges, cuts):
    for left, right in cuts:
        remaining = []
        for start, end in ranges:
            if right < start or left > end:
                remaining.append((start, end))
            else:
                if start < left:
                    remaining.append((start, left - 1))
                if end > right:
                    remaining.append((right + 1, end))
        ranges = remaining
    return ranges


def _intersect(ranges, candidates):
    return [(max(a, c), min(b, d)) for a, b in ranges for c, d in candidates if max(a, c) <= min(b, d)]


def _interval(net):
    return int(net.network_address), int(net.broadcast_address)


def _public_space():
    # Conservative routable-source model; service tags are deliberately not resolved to static IPs.
    reserved4 = "0.0.0.0/8 10.0.0.0/8 100.64.0.0/10 127.0.0.0/8 169.254.0.0/16 172.16.0.0/12 192.0.0.0/24 192.0.2.0/24 192.168.0.0/16 198.18.0.0/15 198.51.100.0/24 203.0.113.0/24 224.0.0.0/4 240.0.0.0/4"
    v4 = _subtract([(0, 2**32 - 1)], [_interval(ipaddress.ip_network(n)) for n in reserved4.split()])
    v4 += [_interval(ipaddress.ip_network(n)) for n in ("192.0.0.9/32", "192.0.0.10/32")]
    v6 = _subtract([_interval(ipaddress.ip_network("2000::/3"))], [_interval(ipaddress.ip_network(n)) for n in ("2001:db8::/32", "2001:2::/48")])
    return {4: v4, 6: v6}


def evaluate_rules(config):
    """First-match CIDR evaluation, with conservative handling of header/tag conditions."""
    rules = config.get("ipSecurityRestrictions") or []
    if not isinstance(rules, list):
        return {"exposure": "Indeterminada", "detail": "Formato de restricciones no reconocido.", "default": "", "ipv4_open": "Indeterminado"}
    default = config.get("ipSecurityRestrictionsDefaultAction") or ("Deny" if rules else "Allow")
    if default not in {"Allow", "Deny"}:
        return {"exposure": "Indeterminada", "detail": "Acción por defecto no reconocida.", "default": str(default), "ipv4_open": "Indeterminado"}
    try:
        ordered = sorted(rules, key=lambda r: int(r.get("priority", 2147483647)))
    except (TypeError, ValueError, AttributeError):
        return {"exposure": "Indeterminada", "detail": "Regla/prioridad no reconocida.", "default": default, "ipv4_open": "Indeterminado"}
    initial = _public_space()
    remaining = {family: list(values) for family, values in initial.items()}
    allowed = {4: [], 6: []}
    possible = False
    conditional = False
    unknown = False
    priorities = {}
    for rule in ordered:
        action = rule.get("action")
        priority = int(rule.get("priority", 2147483647))
        if priority in priorities and priorities[priority] != action:
            unknown = True  # conflicting equal priorities: do not certify an order.
        priorities[priority] = action
        if action not in {"Allow", "Deny"}:
            unknown = True
            continue
        source = str(rule.get("ipAddress") or "").strip()
        if rule.get("vnetSubnetResourceId") or source.lower() == "anyvnets":
            continue  # VNet/service endpoint access is not a public Internet source.
        has_headers = any(rule.get("headers", {}).values()) if isinstance(rule.get("headers"), dict) else bool(rule.get("headers"))
        is_tag = str(rule.get("tag", "")).lower() == "servicetag"
        nets = {4: [], 6: []}
        if source.lower() == "any":
            nets = {4: [(0, 2**32 - 1)], 6: [(0, 2**128 - 1)]}
        elif is_tag:
            nets = {family: list(values) for family, values in remaining.items()}
        else:
            try:
                for item in source.split(","):
                    net = ipaddress.ip_network(item.strip(), strict=False)
                    nets[net.version].append(_interval(net))
            except ValueError:
                unknown = True
                continue
        matches = {family: _intersect(remaining[family], nets[family]) for family in (4, 6)}
        if not any(matches.values()):
            continue
        if has_headers or is_tag:
            conditional = True
            if action == "Allow":
                possible = True
            continue
        for family in (4, 6):
            if action == "Allow":
                allowed[family].extend(matches[family])
            remaining[family] = _subtract(remaining[family], nets[family])
    if default == "Allow":
        for family in (4, 6):
            allowed[family].extend(remaining[family])
    full4 = not _subtract(initial[4], allowed[4])
    full6 = not _subtract(initial[6], allowed[6])
    if unknown:
        return {"exposure": "Indeterminada", "detail": "Regla no interpretable o prioridades iguales con acciones distintas; revisión necesaria.", "default": default, "ipv4_open": "Indeterminado"}
    if (full4 or full6) and not conditional:
        return {"exposure": "Sí", "detail": "Sin restricciones de origen para " + ("IPv4 e IPv6." if full4 and full6 else "IPv4." if full4 else "IPv6."), "default": default, "ipv4_open": "Sí" if full4 else "No"}
    if any(allowed.values()) or possible:
        return {"exposure": "Restringida", "detail": "Endpoint público limitado por CIDR, service tags o condiciones de cabecera; no equivale a acceso privado.", "default": default, "ipv4_open": "No" if not conditional else "Indeterminado"}
    return {"exposure": "No", "detail": "No se identificó una regla efectiva que permita orígenes públicos; pueden existir accesos VNet/privados.", "default": default, "ipv4_open": "No"}


def network_assessment(site, config, config_ok, private_connections, ase):
    props = properties(site)
    pna = props.get("publicNetworkAccess")
    if pna is None:
        pna = config.get("publicNetworkAccess")
    if str(pna).lower() == "disabled":
        return {"exposure": "No", "detail": "publicNetworkAccess=Disabled.", "public_access": "Disabled", "default": "", "ipv4_open": "No"}
    if props.get("hostingEnvironmentProfile"):
        if ase is None:
            return {"exposure": "Indeterminada", "detail": "No se pudo leer si el App Service Environment usa ILB.", "public_access": str(pna or "No establecido"), "default": "", "ipv4_open": "Indeterminado"}
        mode = properties(ase).get("internalLoadBalancingMode")
        if mode in (1, 3) or "web" in str(mode).lower():
            return {"exposure": "No", "detail": "App Service Environment con balanceador interno para Web; un proxy público externo queda fuera de esta evaluación directa.", "public_access": str(pna or "No establecido"), "default": "", "ipv4_open": "No"}
        if mode not in (0, 2, "None", "none", "0", "2", "Publishing", "publishing"):
            return {"exposure": "Indeterminada", "detail": "Modo del App Service Environment no reconocido.", "public_access": str(pna or "No establecido"), "default": "", "ipv4_open": "Indeterminado"}
    if not pna:
        if private_connections is None:
            return {"exposure": "Indeterminada", "detail": "publicNetworkAccess no establecido y conexiones privadas no legibles.", "public_access": "No establecido", "default": "", "ipv4_open": "Indeterminado"}
        if private_connections:
            return {"exposure": "No", "detail": "publicNetworkAccess no establecido y existe private endpoint: acceso público deshabilitado por defecto de App Service.", "public_access": "Disabled (predeterminado con PE)", "default": "", "ipv4_open": "No"}
    elif str(pna).lower() != "enabled":
        return {"exposure": "Indeterminada", "detail": "Valor publicNetworkAccess no reconocido.", "public_access": str(pna), "default": "", "ipv4_open": "Indeterminado"}
    if not config_ok:
        return {"exposure": "Indeterminada", "detail": "No se pudo leer config/web; restricciones de acceso no verificadas.", "public_access": str(pna or "Enabled (predeterminado sin PE)"), "default": "", "ipv4_open": "Indeterminado"}
    result = evaluate_rules(config)
    result["public_access"] = str(pna or "Enabled (predeterminado sin PE)")
    return result


def resource_group(resource_id):
    match = re.search(r"/resourceGroups/([^/]+)", resource_id, re.I)
    return match.group(1) if match else ""


def safe_hostname(value):
    value = str(value or "")
    return value if re.fullmatch(r"[A-Za-z0-9.-]+", value) else ""


def inventory(client, subscriptions, include_slots=True, include_image_hints=False, workers=4):
    errors, php_rows, rule_rows, coverage, targets = [], [], [], [], []

    def error(sub, resource, operation, exc):
        errors.append({"Suscripción": sub, "Recurso": resource, "Operación GET": operation,
                       "Código": exc.code, "Detalle": str(exc)})

    for sub in subscriptions:
        sid = sub["id"]
        before = len(targets)
        complete = True
        try:
            for site in client.pages(f"/subscriptions/{sid}/providers/Microsoft.Web/sites?api-version={API}"):
                if not web_app(site):
                    continue
                targets.append((sub, site))
                if include_slots:
                    try:
                        for slot in client.pages(site["id"] + f"/slots?api-version={API}"):
                            targets.append((sub, slot))
                    except ReadError as exc:
                        complete = False
                        error(sid, site["id"], "Listar slots", exc)
        except ReadError as exc:
            complete = False
            error(sid, "", "Listar Web Apps", exc)
        coverage.append({"Suscripción": sub.get("name", sid), "ID suscripción": sid,
                         "Recursos y slots enumerados": len(targets) - before,
                         "Enumeración completa": "Sí" if complete else "No"})

    def inspect(target):
        sub, site = target
        sid, rid = sub["id"], site["id"]
        details = properties(site)
        config_ok = True
        try:
            config = properties(client.get(rid + f"/config/web?api-version={API}"))
        except ReadError as exc:
            config_ok = False
            config = details.get("siteConfig") or {}
            error(sid, rid, "Leer config/web", exc)
        runtime = php_runtime(config, site.get("kind", ""), include_image_hints)
        if not runtime:
            fx = str(config.get("linuxFxVersion") or config.get("windowsFxVersion") or "")
            if re.match(r"^(DOCKER|COMPOSE|KUBE)\||^sitecontainers$", fx, re.I) or "container" in site.get("kind", "").lower():
                error(sid, rid, "Determinar runtime de contenedor", ReadError("ContainerNotInspected", "No se inspeccionó la imagen ni el proceso del contenedor; PHP puede estar presente sin figurar en ARM."))
            elif config_ok and "linux" in site.get("kind", "").lower() and not fx:
                error(sid, rid, "Determinar runtime Linux", ReadError("RuntimeNotDetermined", "ARM no devuelve linuxFxVersion; no se puede excluir PHP con esta lectura."))
            elif not config_ok and not config:
                pass  # The original GET failure is already recorded; do not claim absence of PHP.
            return None
        try:
            private = list(client.pages(rid + f"/privateEndpointConnections?api-version={API}"))
        except ReadError as exc:
            private = None
            error(sid, rid, "Listar private endpoints", exc)
        ase = None
        ase_id = (details.get("hostingEnvironmentProfile") or {}).get("id")
        if ase_id:
            try:
                ase = client.get(ase_id + f"?api-version={API}")
            except ReadError as exc:
                error(sid, rid, "Leer App Service Environment", exc)
        network = network_assessment(site, config, config_ok, private, ase)
        name = str(site.get("name", ""))
        is_slot = "/slots/" in rid.lower()
        connections = private if private is not None else []
        states = [str((properties(p).get("privateLinkServiceConnectionState") or {}).get("status", "Desconocido")) for p in connections]
        row = {"Suscripción": sub.get("name", sid), "ID suscripción": sid,
               "Grupo de recursos": resource_group(rid), "Web App / slot": name,
               "Tipo": "Slot" if is_slot else "Web App", "Región": site.get("location", ""),
               "Estado": details.get("state", "No disponible"), "Evidencia PHP": runtime["evidence"],
               "Versión PHP configurada": runtime["version"], "Campo de evidencia": runtime["source"],
               "Exposición pública configurada": network["exposure"], "Motivo de exposición": network["detail"],
               "IPv4 sin restricciones de origen": network["ipv4_open"], "publicNetworkAccess efectivo": network["public_access"],
               "Acción ACL no coincidente": network["default"],
               "Private endpoints": len(connections) if private is not None else "No legible",
               "Private endpoints aprobados": sum(s.lower() == "approved" for s in states) if private is not None else "No legible",
               "Estados private endpoint": ", ".join(states),
               "Hostname predeterminado": safe_hostname(details.get("defaultHostName")),
               "Hostnames habilitados": ", ".join(filter(None, (safe_hostname(v) for v in details.get("enabledHostNames", [])))),
               "HTTPS obligatorio": details.get("httpsOnly"),
               "ID recurso": rid,
               "Límite de interpretación": "Configuración ARM; no prueba externa ni inspección de procesos. No confirma versión instalada, parches o vulnerabilidad."}
        # Main application restrictions only. SCM/Kudu is a separate management endpoint.
        for rule in config.get("ipSecurityRestrictions") or []:
            rule_rows.append({"ID recurso": rid, "Web App / slot": name, "Prioridad": rule.get("priority"),
                              "Acción": rule.get("action"), "Origen IP / tag": str(rule.get("ipAddress") or ""),
                              "Tipo tag": rule.get("tag", ""), "Subnet VNet": rule.get("vnetSubnetResourceId", ""),
                              "Cabeceras condicionadas (nombres)": ", ".join(sorted((rule.get("headers") or {}).keys()))})
        return row

    # list.map preserves resource order; CLI invocations use argument lists, never shell strings.
    with ThreadPoolExecutor(max_workers=workers) as pool:
        for index, result in enumerate(pool.map(inspect, targets), 1):
            if result:
                php_rows.append(result)
            if index % 20 == 0 or index == len(targets):
                print(f"Revisados {index}/{len(targets)} recursos; coincidencias PHP: {len(php_rows)}", file=sys.stderr)
    for c in coverage:
        sid = c["ID suscripción"]
        c["Recursos PHP incluidos"] = sum(r["ID suscripción"] == sid for r in php_rows)
        c["Incidencias / límites"] = sum(e["Suscripción"] == sid for e in errors)
    return {"php": sorted(php_rows, key=lambda r: r["ID recurso"].lower()),
            "rules": sorted(rule_rows, key=lambda r: (r["ID recurso"], int(r["Prioridad"] or 0))),
            "errors": sorted(errors, key=lambda e: (e["Suscripción"], e["Recurso"], e["Operación GET"])),
            "coverage": coverage, "scanned": len(targets), "include_slots": include_slots,
            "include_image_hints": include_image_hints, "checked_utc": datetime.now(timezone.utc).isoformat()}


def write_excel(data, output):
    # Runtime dependency for the user's Azure/Cloud Shell environment; no Azure SDK is needed.
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:
        raise ReadError("MissingDependency", "Instala las dependencias: python -m pip install -r requirements.txt") from exc
    wb = Workbook()
    summary = wb.active
    summary.title = "Resumen"
    summary.append(["Inventario Azure Web Apps con PHP"])
    summary.append(["Fecha UTC", data["checked_utc"]])
    summary.append(["Recursos y slots revisados", data["scanned"]])
    summary.append(["Web Apps / slots PHP incluidos", len(data["php"])])
    for label in ("Sí", "Restringida", "No", "Indeterminada"):
        summary.append(["Exposición pública: " + label, sum(r["Exposición pública configurada"] == label for r in data["php"])])
    summary.append(["Incidencias y límites", len(data["errors"])])
    summary.append(["Slots incluidos", "Sí" if data["include_slots"] else "No"])
    summary.append(["Imágenes con indicio PHP incluidas", "Sí (indicios separados)" if data["include_image_hints"] else "No"])
    summary.append(["Cobertura", "PARCIAL: consulta Incidencias y Cobertura" if data["errors"] else "Lecturas completadas para el alcance seleccionado"])
    notes = [
        ("Global Administrator", "El rol de Entra no concede automáticamente lectura de recursos Azure. Reader de Azure RBAC en las suscripciones es suficiente para las operaciones previstas."),
        ("Filtro PHP", "PHP configurado en linuxFxVersion/windowsFxVersion/phpVersion. No se verifica un proceso PHP ni el contenido de imágenes personalizadas."),
        ("Sí", "El endpoint predeterminado permite orígenes públicos sin restricciones en al menos una familia IP."),
        ("Restringida", "El endpoint es público con condiciones CIDR/service tags/cabeceras. No debe clasificarse como privado."),
        ("No", "No se identifica acceso público directo por la configuración: PNA deshabilitado, ASE interno o ACL sin orígenes públicos permitidos."),
        ("Indeterminada", "Faltan permisos/datos o hay reglas que no se pueden evaluar de forma fiable."),
        ("Private endpoint", "Su existencia no desactiva por sí sola un publicNetworkAccess explícitamente Enabled. Con PNA no establecido se aplica el comportamiento predeterminado documentado."),
        ("Alcance de red", "No se prueba conectividad desde Internet; no se evalúan rutas indirectas mediante Front Door, Application Gateway, APIM, proxies u otros recursos."),
        ("Límites ACL", "Evaluación conservadora de IP públicas, prioridades y acción por defecto. No resuelve rangos actuales de service tags ni demuestra que una cabecera sea de un proxy de confianza."),
        ("SCM y autenticación", "Se evalúa la aplicación principal. SCM/Kudu, Easy Auth, autorizaciones de la aplicación y WAF no forman parte de esta clasificación de red."),
        ("Estado de ejecución", "Los recursos detenidos se conservan con su configuración. Una app detenida no sirve tráfico aunque tenga acceso público configurado."),
        ("CVE PHP", "La versión configurada puede ser solo una rama (8.2). No se determina CVE-2026-109618/109619 sin versión instalada y evidencia de parche del proveedor."),
        ("Documentación de red", NETWORK_DOC),
        ("API config/web", "https://learn.microsoft.com/rest/api/appservice/web-apps/get-configuration?view=rest-appservice-2024-11-01"),
    ]
    summary.append([])
    for item in notes:
        summary.append(list(item))
    headers = {"WebApps PHP": ["Suscripción", "ID suscripción", "Grupo de recursos", "Web App / slot", "Tipo", "Región", "Estado", "Evidencia PHP", "Versión PHP configurada", "Campo de evidencia", "Exposición pública configurada", "Motivo de exposición", "IPv4 sin restricciones de origen", "publicNetworkAccess efectivo", "Acción ACL no coincidente", "Private endpoints", "Private endpoints aprobados", "Estados private endpoint", "Hostname predeterminado", "Hostnames habilitados", "HTTPS obligatorio", "ID recurso", "Límite de interpretación"],
               "Reglas acceso": ["ID recurso", "Web App / slot", "Prioridad", "Acción", "Origen IP / tag", "Tipo tag", "Subnet VNet", "Cabeceras condicionadas (nombres)"],
               "Incidencias": ["Suscripción", "Recurso", "Operación GET", "Código", "Detalle"],
               "Cobertura": ["Suscripción", "ID suscripción", "Recursos y slots enumerados", "Enumeración completa", "Recursos PHP incluidos", "Incidencias / límites"]}
    for number, (name, records) in enumerate(zip(headers, (data["php"], data["rules"], data["errors"], data["coverage"])), 1):
        ws = wb.create_sheet(name)
        ws.append(headers[name])
        for record in records:
            ws.append([record.get(k) for k in headers[name]])
        ws.freeze_panes = "E2" if name == "WebApps PHP" else "A2"
        if records:
            table = Table(displayName=f"InventoryTable{number}", ref=ws.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
        ws.auto_filter.ref = ws.dimensions
    for ws in wb:
        ws.sheet_view.showGridLines = False
        for row in ws:
            for cell in row:
                # Explicit strings prevent formula injection from resource names or metadata.
                if isinstance(cell.value, str):
                    cell.value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", cell.value)[:32767]
                    cell.data_type = "s"
                cell.font = Font(name="Calibri", size=11)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="243C5A")
            cell.font = Font(name="Calibri", size=11, color="FFFFFF", bold=True)
        ws.row_dimensions[1].height = 30
        for col in ws.columns:
            letter = col[0].column_letter
            width = max((len(str(c.value or "")) for c in col), default=12)
            ws.column_dimensions[letter].width = min(max(width + 2, 16), 65)
    summary.column_dimensions["A"].width = 40
    summary.column_dimensions["B"].width = 112
    for row in range(14, summary.max_row + 1):
        summary.row_dimensions[row].height = 42
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        raise ReadError("OutputExists", "El archivo de salida ya existe; elige otro nombre para conservar el informe anterior.")
    try:
        with output.open("xb") as stream:
            wb.save(stream)
    except FileExistsError as exc:
        raise ReadError("OutputExists", "Otra ejecución creó el archivo; no se sobrescribió.") from exc


def main(argv=None):
    parser = argparse.ArgumentParser(description="Inventario de solo lectura de Web Apps/slots con PHP y exposición de red configurada. No requiere Global Administrator; sí Azure RBAC Reader en el alcance.")
    parser.add_argument("--subscriptions", nargs="+", help="IDs de suscripción. Por defecto: todas las Enabled visibles en az account list --all.")
    parser.add_argument("--output", default="azure_webapps_php.xlsx", help="Excel nuevo; no se sobrescriben archivos existentes.")
    parser.add_argument("--no-slots", action="store_true", help="Omitir slots de despliegue.")
    parser.add_argument("--include-container-hints", action="store_true", help="Incluir nombres de imagen que sugieren PHP, claramente marcados como NO confirmados.")
    parser.add_argument("--workers", type=int, default=4, metavar="N", help="Concurrencia de lectura: 1 a 16 (por defecto: 4).")
    parser.add_argument("--timeout", type=int, default=60, metavar="SEGUNDOS", help="Tiempo máximo por comando: 10 a 300 (por defecto: 60).")
    args = parser.parse_args(argv)
    if not 1 <= args.workers <= 16 or not 10 <= args.timeout <= 300:
        parser.error("--workers debe estar entre 1 y 16; --timeout, entre 10 y 300.")
    if Path(args.output).exists():
        parser.error("El archivo de salida existe; utiliza otro nombre.")
    try:
        if args.subscriptions:
            for value in args.subscriptions:
                UUID(value)
        client = AzureCLI(args.timeout)
        visible = client.subscriptions()
        names = {s["id"].lower(): s for s in visible}
        if args.subscriptions:
            subscriptions = [names.get(s.lower(), {"id": s, "name": s}) for s in dict.fromkeys(args.subscriptions)]
        else:
            subscriptions = list({s["id"].lower(): s for s in visible if s.get("state") == "Enabled"}.values())
        if not subscriptions:
            raise ReadError("NoSubscriptions", "No hay suscripciones visibles. Verifica tenant, sesión az login, PIM y Azure RBAC Reader; Global Administrator de Entra no basta por sí solo.")
        print(f"Solo lectura: {len(subscriptions)} suscripciones. No se accede a secretos ni se cambia Azure.", file=sys.stderr)
        data = inventory(client, subscriptions, not args.no_slots, args.include_container_hints, args.workers)
        write_excel(data, args.output)
        print(f"Excel: {Path(args.output).resolve()}")
        print(f"PHP: {len(data['php'])}; recursos revisados: {data['scanned']}; incidencias/límites: {len(data['errors'])}.")
        return 2 if data["errors"] else 0
    except ValueError:
        parser.error("Los IDs de suscripción deben ser UUID válidos.")
    except ReadError as exc:
        print(f"{exc.code}: {exc}", file=sys.stderr)
        return 1
    return 1


if __name__ == "__main__":
    sys.exit(main())
